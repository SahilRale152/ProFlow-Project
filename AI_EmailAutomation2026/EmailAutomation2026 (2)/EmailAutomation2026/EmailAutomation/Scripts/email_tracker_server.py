"""
============================================================================
EMAIL OPEN TRACKING SERVER
============================================================================
Purpose:
  - Serves a 1x1 invisible tracking pixel in each email
  - When recipient opens the email, their client loads the pixel from this server
  - Server logs the open event and updates Email_Tracking.xlsx automatically

HOW IT WORKS:
  1. Each email gets a unique tracking URL like:
       http://YOUR_SERVER_IP:5050/track?email=ceo@company.com&company=CompanyName
  2. When email is opened, that URL is hit → this server responds with a tiny image
  3. Server writes "Opened" + timestamp into Email_Tracking.xlsx

SETUP:
  pip install flask openpyxl

RUN:
  python email_tracker_server.py

IMPORTANT:
  - This server must be running and accessible from the internet
  - Use ngrok (free) to expose localhost: ngrok http 5050
  - Set TRACKING_BASE_URL in Config (notebook) to your ngrok URL
============================================================================
"""

from flask import Flask, request, send_file, make_response
import openpyxl
from openpyxl.styles import PatternFill, Font
from datetime import datetime
import io
import base64
import os
import threading
import logging

app = Flask(__name__)

# ─── CONFIGURE THESE ──────────────────────────────────────────────────────────
# Resolved relative to this script's own location, not hardcoded to one
# person's folder -- see RUN_CAMPAIGN.py Config for why.
from pathlib import Path as _Path
_DATA_DIR = _Path(__file__).resolve().parent.parent / "Data"
TRACKER_FILE = str(_DATA_DIR / "Email_Tracking_NEW.xlsx")  # Same path as in your notebook Config
LOG_FILE     = "open_tracking.log"
# ──────────────────────────────────────────────────────────────────────────────

# Thread lock so concurrent opens don't corrupt the Excel file
excel_lock = threading.Lock()

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s - %(message)s"
)

# Tiny 1×1 transparent GIF in base64
PIXEL_B64 = (
    "R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBRAA7"
)
PIXEL_BYTES = base64.b64decode(PIXEL_B64)


def update_excel_open_status(email_address: str, company: str):
    """
    Open Email_Tracking.xlsx, find the row matching the email address,
    and set 'Email Open Status' = 'Opened' with a timestamp.
    If the column doesn't exist yet, create it.
    """
    if not os.path.exists(TRACKER_FILE):
        logging.warning(f"Tracker file not found: {TRACKER_FILE}")
        return

    with excel_lock:
        try:
            wb = openpyxl.load_workbook(TRACKER_FILE)
            # Never use wb.active — it may be the 'Summary' sheet (no Email
            # column). Prefer the master tracking sheet.
            if "Total Sent" in wb.sheetnames:
                ws = wb["Total Sent"]
            elif "Current Sent" in wb.sheetnames:
                ws = wb["Current Sent"]
            elif "Sent" in wb.sheetnames:
                ws = wb["Sent"]
            else:
                ws = wb.active

            # ── Find or create the tracking columns ──────────────────────────
            headers = [cell.value for cell in ws[1]]

            def get_or_create_col(name):
                if name in headers:
                    return headers.index(name) + 1          # 1-based
                else:
                    col_idx = len(headers) + 1
                    ws.cell(row=1, column=col_idx, value=name).font = Font(bold=True)
                    headers.append(name)
                    return col_idx

            status_col    = get_or_create_col("Email Open Status")
            open_time_col = get_or_create_col("Email Opened At")
            open_count_col = get_or_create_col("Open Count")

            email_col_idx = None
            for idx, h in enumerate(headers, 1):
                if h and str(h).strip().lower() == "email":
                    email_col_idx = idx
                    break

            if email_col_idx is None:
                logging.error("No 'Email' column found in tracker file.")
                return

            # ── Walk rows, match email, update ───────────────────────────────
            matched = False
            now_str = datetime.now().strftime("%d-%m-%Y %H:%M:%S")

            for row in ws.iter_rows(min_row=2):
                cell_email = row[email_col_idx - 1].value
                if cell_email and str(cell_email).strip().lower() == email_address.lower():
                    matched = True
                    row_num = row[0].row

                    # Update status
                    status_cell = ws.cell(row=row_num, column=status_col)
                    status_cell.value = "Opened"
                    status_cell.fill  = PatternFill("solid", fgColor="C6EFCE")   # green
                    status_cell.font  = Font(color="276221", bold=True)

                    # Set first-open timestamp (don't overwrite if already set)
                    time_cell = ws.cell(row=row_num, column=open_time_col)
                    if not time_cell.value:
                        time_cell.value = now_str

                    # Increment open count
                    count_cell = ws.cell(row=row_num, column=open_count_col)
                    count_cell.value = (int(count_cell.value or 0)) + 1

                    logging.info(f"OPENED | {email_address} | {company} | {now_str} | count={count_cell.value}")
                    print(f"[OPEN TRACKED] {email_address} ({company}) at {now_str}")

            if not matched:
                logging.warning(f"Email not found in tracker: {email_address}")

            wb.save(TRACKER_FILE)

        except Exception as e:
            logging.error(f"Excel update error: {e}")
            print(f"[ERROR] Could not update Excel: {e}")


@app.route("/track")
def track_open():
    """
    Tracking pixel endpoint.
    URL params:
      ?email=recipient@email.com&company=CompanyName
    Returns a 1×1 transparent GIF.
    """
    email_param   = request.args.get("email", "").strip()
    company_param = request.args.get("company", "").strip()

    if email_param:
        # Run update in background so the image loads instantly
        threading.Thread(
            target=update_excel_open_status,
            args=(email_param, company_param),
            daemon=True
        ).start()

    response = make_response(PIXEL_BYTES)
    response.headers["Content-Type"]  = "image/gif"
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"]        = "no-cache"
    return response


@app.route("/status")
def status():
    """Quick health-check — visit this in browser to confirm server is running."""
    return {"status": "running", "tracker_file": TRACKER_FILE}, 200


if __name__ == "__main__":
    print("=" * 60)
    print("  EMAIL OPEN TRACKING SERVER")
    print("=" * 60)
    print(f"  Tracker file : {TRACKER_FILE}")
    print(f"  Log file     : {LOG_FILE}")
    print()
    print("  Server running at: http://localhost:5050")
    print()
    print("  To expose to internet (required for tracking):")
    print("    1. Install ngrok: https://ngrok.com/download")
    print("    2. Run: ngrok http 5050")
    print("    3. Copy the https://xxxx.ngrok.io URL")
    print("    4. Set TRACKING_BASE_URL in your notebook Config")
    print("=" * 60)
    app.run(host="0.0.0.0", port=5050, debug=False)
